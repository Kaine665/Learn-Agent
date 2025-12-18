"""
通用文件转换提取框架
实现智能的、可学习的文件转换系统
"""

import json
import os
import re
from typing import Dict, List, Any, Optional, Callable, Tuple
from datetime import datetime
from dataclasses import dataclass, asdict
from models import Transaction


@dataclass
class FormatInfo:
    """格式信息"""
    file_type: str  # xlsx, csv, json
    source_type: str  # wechat, alipay, bank, unknown
    structure: Dict[str, Any]  # 结构信息
    signatures: List[str]  # 格式特征签名


@dataclass
class ExtractionResult:
    """提取结果"""
    success: bool
    confidence: float  # 0.0-1.0
    data: List[Transaction]
    issues: List[str]  # 无法确定的问题
    validation_errors: List[str]


@dataclass
class ConversionRule:
    """转换规则"""
    rule_id: str
    source_type: str
    file_format: str
    format_signatures: List[str]
    structure_mapping: Dict[str, Any]
    transformation_logic: Dict[str, Any]
    validation_rules: List[str]
    created_at: str
    updated_at: str
    usage_count: int = 0
    success_rate: float = 1.0
    
    def to_dict(self) -> Dict[str, Any]:
        """转换为字典"""
        return asdict(self)
    
    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> 'ConversionRule':
        """从字典创建"""
        return cls(**data)


class FormatDetector:
    """格式检测器"""
    
    def detect(self, file_path: str) -> FormatInfo:
        """
        检测文件格式
        
        Returns:
            FormatInfo: 格式信息
        """
        # 检测文件扩展名
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext == '.xlsx' or file_ext == '.xls':
            return self._detect_excel(file_path)
        elif file_ext == '.csv':
            return self._detect_csv(file_path)
        elif file_ext == '.json':
            return self._detect_json(file_path)
        else:
            return FormatInfo(
                file_type=file_ext[1:] if file_ext else "unknown",
                source_type="unknown",
                structure={},
                signatures=[]
            )
    
    def _detect_excel(self, file_path: str) -> FormatInfo:
        """检测Excel文件格式"""
        try:
            from openpyxl import load_workbook
            
            wb = load_workbook(file_path)
            ws = wb.active
            
            signatures = []
            metadata_rows = []
            header_row = None
            data_start_row = None
            
            # 扫描前50行，查找特征
            for row_idx in range(1, min(50, ws.max_row + 1)):
                row = [cell.value for cell in ws[row_idx]]
                row_str = ' '.join([str(cell).strip() for cell in row if cell]).lower()
                
                # 检测微信支付账单特征
                if '微信支付账单明细' in row_str:
                    signatures.append("包含'微信支付账单明细'")
                    source_type = "wechat"
                elif '支付宝' in row_str or 'alipay' in row_str:
                    signatures.append("包含'支付宝'")
                    source_type = "alipay"
                elif '银行' in row_str or 'bank' in row_str:
                    signatures.append("包含'银行'")
                    source_type = "bank"
                
                # 检测元数据行
                if any(keyword in row_str for keyword in ['起始时间', '终止时间', '导出类型', '导出时间']):
                    metadata_rows.append(row_idx)
                
                # 检测表头行（包含"交易时间"和"金额"）
                if '交易时间' in row_str and ('金额' in row_str or '收/支' in row_str):
                    if header_row is None:
                        header_row = row_idx
                        data_start_row = row_idx + 1
                        signatures.append(f"表头在第{row_idx}行")
            
            # 检测列特征
            if header_row:
                header_row_data = [cell.value for cell in ws[header_row]]
                header_str = ' '.join([str(cell).strip() for cell in header_row_data if cell])
                
                if '收/支' in header_str:
                    signatures.append("有'收/支'列")
                if '交易类型' in header_str:
                    signatures.append("有'交易类型'列")
                if '商品' in header_str:
                    signatures.append("有'商品'列")
                if '交易对方' in header_str:
                    signatures.append("有'交易对方'列")
            
            return FormatInfo(
                file_type="xlsx",
                source_type=source_type if 'source_type' in locals() else "unknown",
                structure={
                    "metadata_rows": metadata_rows,
                    "header_row": header_row,
                    "data_start_row": data_start_row
                },
                signatures=signatures
            )
        except Exception as e:
            return FormatInfo(
                file_type="xlsx",
                source_type="unknown",
                structure={},
                signatures=[f"检测失败: {str(e)}"]
            )
    
    def _detect_csv(self, file_path: str) -> FormatInfo:
        """检测CSV文件格式"""
        import csv
        signatures = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                first_row = next(reader, None)
                
                if first_row:
                    header_str = ' '.join(first_row).lower()
                    if 'date' in header_str or '日期' in header_str:
                        signatures.append("有日期列")
                    if 'amount' in header_str or '金额' in header_str:
                        signatures.append("有金额列")
        except:
            pass
        
        return FormatInfo(
            file_type="csv",
            source_type="unknown",
            structure={},
            signatures=signatures
        )
    
    def _detect_json(self, file_path: str) -> FormatInfo:
        """检测JSON文件格式"""
        signatures = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                if isinstance(data, list):
                    signatures.append("JSON数组格式")
                elif isinstance(data, dict):
                    signatures.append("JSON对象格式")
                    if 'transactions' in data:
                        signatures.append("包含'transactions'字段")
        except:
            pass
        
        return FormatInfo(
            file_type="json",
            source_type="unknown",
            structure={},
            signatures=signatures
        )
    
    def extract_signatures(self, file_path: str) -> List[str]:
        """提取格式特征签名，用于规则匹配"""
        format_info = self.detect(file_path)
        return format_info.signatures


class RuleManager:
    """规则管理器"""
    
    def __init__(self, rules_dir: str = "data/rules"):
        """
        Args:
            rules_dir: 规则存储目录
        """
        self.rules_dir = rules_dir
        os.makedirs(rules_dir, exist_ok=True)
        self.rules: Dict[str, ConversionRule] = {}
        self._load_rules()
    
    def _load_rules(self):
        """加载所有规则"""
        if not os.path.exists(self.rules_dir):
            return
        
        for filename in os.listdir(self.rules_dir):
            if filename.endswith('.json'):
                rule_path = os.path.join(self.rules_dir, filename)
                try:
                    with open(rule_path, 'r', encoding='utf-8') as f:
                        rule_data = json.load(f)
                        rule = ConversionRule.from_dict(rule_data)
                        self.rules[rule.rule_id] = rule
                except Exception as e:
                    print(f"⚠️ 加载规则失败 {filename}: {e}")
    
    def find_rule(self, format_info: FormatInfo) -> Optional[ConversionRule]:
        """
        查找匹配的规则
        
        Args:
            format_info: 格式信息
        
        Returns:
            匹配的规则，如果没有则返回None
        """
        file_signatures = format_info.signatures
        
        best_match = None
        best_score = 0.0
        
        for rule in self.rules.values():
            # 检查文件格式和源类型
            if rule.file_format != format_info.file_type:
                continue
            
            if rule.source_type != format_info.source_type and format_info.source_type != "unknown":
                continue
            
            # 匹配格式特征
            score = self._match_signatures(file_signatures, rule.format_signatures)
            
            if score > best_score and score >= 0.8:  # 匹配度阈值
                best_score = score
                best_match = rule
        
        return best_match
    
    def _match_signatures(self, file_signatures: List[str], rule_signatures: List[str]) -> float:
        """
        匹配格式特征，返回匹配度（0.0-1.0）
        
        Args:
            file_signatures: 文件特征
            rule_signatures: 规则特征
        
        Returns:
            匹配度
        """
        if not rule_signatures:
            return 0.0
        
        matched = 0
        for rule_sig in rule_signatures:
            for file_sig in file_signatures:
                if rule_sig.lower() in file_sig.lower() or file_sig.lower() in rule_sig.lower():
                    matched += 1
                    break
        
        return matched / len(rule_signatures)
    
    def save_rule(self, rule: ConversionRule):
        """保存规则"""
        rule.updated_at = datetime.now().isoformat()
        rule.usage_count += 1
        
        rule_path = os.path.join(self.rules_dir, f"{rule.rule_id}.json")
        with open(rule_path, 'w', encoding='utf-8') as f:
            json.dump(rule.to_dict(), f, ensure_ascii=False, indent=2)
        
        self.rules[rule.rule_id] = rule
        print(f"✅ 已保存转换规则：{rule.rule_id}")
    
    def get_rule(self, rule_id: str) -> Optional[ConversionRule]:
        """获取规则"""
        return self.rules.get(rule_id)


class DataExtractor:
    """数据提取器"""
    
    def __init__(self, rule_manager: RuleManager):
        self.rule_manager = rule_manager
    
    def extract(self, file_path: str, rule: ConversionRule) -> ExtractionResult:
        """
        根据规则提取数据
        
        Args:
            file_path: 文件路径
            rule: 转换规则
        
        Returns:
            提取结果
        """
        try:
            if rule.file_format == "xlsx":
                return self._extract_excel(file_path, rule)
            elif rule.file_format == "csv":
                return self._extract_csv(file_path, rule)
            elif rule.file_format == "json":
                return self._extract_json(file_path, rule)
            else:
                return ExtractionResult(
                    success=False,
                    confidence=0.0,
                    data=[],
                    issues=[f"不支持的文件格式：{rule.file_format}"],
                    validation_errors=[]
                )
        except Exception as e:
            return ExtractionResult(
                success=False,
                confidence=0.0,
                data=[],
                issues=[f"提取失败：{str(e)}"],
                validation_errors=[]
            )
    
    def auto_extract(self, file_path: str, format_info: FormatInfo) -> ExtractionResult:
        """
        尝试自动提取（无规则时）
        
        Args:
            file_path: 文件路径
            format_info: 格式信息
        
        Returns:
            提取结果（可能包含问题）
        """
        issues = []
        
        if format_info.file_type == "xlsx":
            # 尝试自动提取Excel
            result = self._auto_extract_excel(file_path, format_info)
            if not result.success:
                issues.extend(result.issues)
            return result
        else:
            return ExtractionResult(
                success=False,
                confidence=0.0,
                data=[],
                issues=["无法自动提取，需要转换规则"],
                validation_errors=[]
            )
    
    def _extract_excel(self, file_path: str, rule: ConversionRule) -> ExtractionResult:
        """根据规则提取Excel数据"""
        from openpyxl import load_workbook
        
        wb = load_workbook(file_path)
        ws = wb.active
        
        structure = rule.structure_mapping
        header_row_num = structure.get("header_row", {}).get("row", None)
        column_mapping = structure.get("column_mapping", {})
        
        if header_row_num is None:
            return ExtractionResult(
                success=False,
                confidence=0.0,
                data=[],
                issues=["规则中未指定表头行"],
                validation_errors=[]
            )
        
        # 读取表头行，建立列索引映射
        header_row = [cell.value for cell in ws[header_row_num]]
        column_indices = {}
        for field, mapping_config in column_mapping.items():
            col_idx = self._find_column_index(header_row, mapping_config)
            if col_idx is not None:
                column_indices[field] = col_idx
        
        if "date" not in column_indices or "amount" not in column_indices:
            return ExtractionResult(
                success=False,
                confidence=0.0,
                data=[],
                issues=["无法找到必要的列（日期或金额）"],
                validation_errors=[]
            )
        
        transactions = []
        validation_errors = []
        
        # 读取数据行
        for row_idx in range(header_row_num + 1, ws.max_row + 1):
            row = [cell.value for cell in ws[row_idx]]
            
            if not any(row):
                continue
            
            try:
                transaction = self._parse_row_with_rule(row, column_indices, column_mapping, rule)
                if transaction:
                    transactions.append(transaction)
            except Exception as e:
                validation_errors.append(f"第{row_idx}行解析失败：{str(e)}")
        
        return ExtractionResult(
            success=True,
            confidence=1.0,
            data=transactions,
            issues=[],
            validation_errors=validation_errors
        )
    
    def _parse_row_with_rule(self, row: List, column_indices: Dict[str, int], column_mapping: Dict, rule: ConversionRule) -> Optional[Transaction]:
        """根据规则解析行数据"""
        # 获取列索引
        date_col = column_indices.get("date")
        amount_col = column_indices.get("amount")
        
        if date_col is None or amount_col is None:
            return None
        
        # 解析日期
        date_value = row[date_col] if date_col < len(row) else None
        date_str = self._parse_date(date_value)
        
        # 解析金额
        amount_value = row[amount_col] if amount_col < len(row) else None
        amount = self._parse_amount(amount_value)
        
        # 判断收入/支出
        is_income = self._determine_income_expense(row, column_indices, column_mapping, rule)
        
        if not is_income and amount > 0:
            amount = -abs(amount)
        
        # 获取类别和描述
        category = self._get_category(row, column_indices, rule, is_income)
        description = self._get_description(row, column_indices)
        
        return Transaction(
            date=date_str,
            amount=amount,
            category=category,
            description=description
        )
    
    def _find_column_index(self, header_row: List, mapping_config: Dict) -> Optional[int]:
        """根据映射配置查找列索引"""
        if not mapping_config:
            return None
        
        patterns = mapping_config.get("patterns", [])
        for col_idx, cell_value in enumerate(header_row):
            if not cell_value:
                continue
            cell_str = str(cell_value).strip()
            for pattern in patterns:
                if pattern in cell_str:
                    return col_idx
        return None
    
    def _parse_date(self, value: Any) -> str:
        """解析日期"""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        elif isinstance(value, str):
            if ' ' in value:
                return value.split(' ')[0]
            return value
        return ""
    
    def _parse_amount(self, value: Any) -> float:
        """解析金额"""
        if isinstance(value, (int, float)):
            return float(value)
        elif isinstance(value, str):
            amount_str = value.replace(',', '').replace('，', '').replace(' ', '').strip()
            amount_str = re.sub(r'[^\d.-]', '', amount_str)
            return float(amount_str) if amount_str else 0.0
        return 0.0
    
    def _determine_income_expense(self, row: List, column_indices: Dict[str, int], column_mapping: Dict, rule: ConversionRule) -> bool:
        """判断收入/支出"""
        logic = rule.transformation_logic.get("is_income", {})
        priority = logic.get("priority", ["income_expense", "type", "description"])
        
        for source in priority:
            if source == "income_expense" and "income_expense" in column_indices:
                col_idx = column_indices["income_expense"]
                if col_idx < len(row) and row[col_idx]:
                    value = str(row[col_idx]).strip()
                    if value == '收' or '收入' in value:
                        return True
                    elif value == '支' or '支出' in value:
                        return False
            
            elif source == "type" and "type" in column_indices:
                col_idx = column_indices["type"]
                if col_idx < len(row) and row[col_idx]:
                    value = str(row[col_idx]).strip()
                    if '收入' in value:
                        return True
                    elif '支出' in value or '支付' in value:
                        return False
        
        # 默认当作支出
        return False
    
    def _get_category(self, row: List, column_indices: Dict[str, int], rule: ConversionRule, is_income: bool) -> str:
        """获取类别"""
        if "category" in column_indices:
            col_idx = column_indices["category"]
            if col_idx < len(row) and row[col_idx]:
                return str(row[col_idx]).strip()
        return "收入" if is_income else "其他"
    
    def _get_description(self, row: List, column_indices: Dict[str, int]) -> str:
        """获取描述"""
        parts = []
        
        if "product" in column_indices:
            product_col = column_indices["product"]
            if product_col < len(row) and row[product_col]:
                parts.append(str(row[product_col]).strip())
        
        if "counterparty" in column_indices:
            counterparty_col = column_indices["counterparty"]
            if counterparty_col < len(row) and row[counterparty_col]:
                counterparty = str(row[counterparty_col]).strip()
                if counterparty not in parts:
                    parts.append(counterparty)
        
        return ' - '.join(parts) if parts else ""
    
    def _auto_extract_excel(self, file_path: str, format_info: FormatInfo) -> ExtractionResult:
        """自动提取Excel（简化版，用于测试）"""
        issues = []
        
        # 这里可以实现自动提取逻辑
        # 但通常需要规则才能准确提取
        
        return ExtractionResult(
            success=False,
            confidence=0.3,
            data=[],
            issues=["无法自动提取，需要转换规则"],
            validation_errors=[]
        )
    
    def _extract_csv(self, file_path: str, rule: ConversionRule) -> ExtractionResult:
        """提取CSV（待实现）"""
        return ExtractionResult(
            success=False,
            confidence=0.0,
            data=[],
            issues=["CSV提取功能待实现"],
            validation_errors=[]
        )
    
    def _extract_json(self, file_path: str, rule: ConversionRule) -> ExtractionResult:
        """提取JSON（待实现）"""
        return ExtractionResult(
            success=False,
            confidence=0.0,
            data=[],
            issues=["JSON提取功能待实现"],
            validation_errors=[]
        )


class FileConverter:
    """文件转换器（主入口）"""
    
    def __init__(self, rules_dir: str = "data/rules"):
        """
        Args:
            rules_dir: 规则存储目录
        """
        self.format_detector = FormatDetector()
        self.rule_manager = RuleManager(rules_dir)
        self.data_extractor = DataExtractor(self.rule_manager)
        # 使用 ConverterAgent 来查找规则（支持快速通道和规则库）
        from converter_agent import ConverterAgent
        self.converter_agent = ConverterAgent(rules_dir)
    
    def convert(self, file_path: str) -> Tuple[bool, List[Transaction], Optional[str]]:
        """
        转换文件（只支持硬编码和规则库匹配）
        
        Args:
            file_path: 文件路径
        
        Returns:
            (success, transactions, message)
        """
        # 1. 使用 ConverterAgent 查找规则（支持快速通道和规则库）
        rule = self.converter_agent.find_rule(file_path)
        
        if rule:
            # 2. 应用规则提取
            print(f"✅ 找到匹配规则：{rule.rule_id}")
            result = self.data_extractor.extract(file_path, rule)
            
            if result.success:
                return True, result.data, f"成功提取 {len(result.data)} 条记录"
            else:
                print(f"❌ 提取失败：{', '.join(result.issues)}")
                return False, [], f"提取失败：{', '.join(result.issues)}"
        else:
            # 3. 找不到规则，已经打印失败信息
            return False, [], "未找到匹配的转换规则"

